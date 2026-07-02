#!/usr/bin/env python3
"""
Extract content and layout metadata from a PowerPoint file (.pptx).

Usage:
    python scripts/extract-pptx.py <input.pptx> [output_dir]

Output:
    <output_dir>/extracted-slides.json
    <output_dir>/assets/slideN_imgM.<ext>

Requires:
    python -m pip install python-pptx

The JSON keeps the original list-of-slides shape used by earlier versions, while
adding richer per-slide metadata for faithful HTML reconstruction.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any


EMU_PER_INCH = 914400
PX_PER_INCH = 96


def load_pptx_modules():
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    except ModuleNotFoundError as exc:
        print(
            "Missing dependency: python-pptx\n\n"
            "Install it with:\n"
            "  python -m pip install python-pptx\n\n"
            "Then rerun:\n"
            "  python scripts/extract-pptx.py <input.pptx> [output_dir]",
            file=sys.stderr,
        )
        raise SystemExit(2) from exc

    return Presentation, MSO_SHAPE_TYPE


def emu_to_inches(value: int | None) -> float | None:
    if value is None:
        return None
    return round(int(value) / EMU_PER_INCH, 4)


def emu_to_px(value: int | None) -> int | None:
    if value is None:
        return None
    return round(int(value) / EMU_PER_INCH * PX_PER_INCH)


def enum_name(value: Any) -> str:
    name = getattr(value, "name", None)
    if name:
        return str(name)
    return str(value)


def geometry(shape: Any) -> dict[str, Any]:
    left = getattr(shape, "left", None)
    top = getattr(shape, "top", None)
    width = getattr(shape, "width", None)
    height = getattr(shape, "height", None)

    return {
        "left": int(left) if left is not None else None,
        "top": int(top) if top is not None else None,
        "width": int(width) if width is not None else None,
        "height": int(height) if height is not None else None,
        "left_px": emu_to_px(left),
        "top_px": emu_to_px(top),
        "width_px": emu_to_px(width),
        "height_px": emu_to_px(height),
        "left_in": emu_to_inches(left),
        "top_in": emu_to_inches(top),
        "width_in": emu_to_inches(width),
        "height_in": emu_to_inches(height),
    }


def shape_alt_text(shape: Any) -> str:
    """Best-effort alt-text extraction from PowerPoint XML."""
    try:
        nodes = shape.element.xpath('.//*[local-name()="cNvPr"]')
    except Exception:
        nodes = []

    for node in nodes:
        descr = node.get("descr")
        title = node.get("title")
        if descr:
            return descr
        if title:
            return title

    return ""


def paragraphs_from_text_frame(shape: Any) -> list[dict[str, Any]]:
    if not getattr(shape, "has_text_frame", False):
        return []

    paragraphs: list[dict[str, Any]] = []
    for paragraph in shape.text_frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            runs.append(
                {
                    "text": run.text,
                    "bold": bool(run.font.bold) if run.font.bold is not None else None,
                    "italic": bool(run.font.italic) if run.font.italic is not None else None,
                    "font_size": int(run.font.size) if run.font.size is not None else None,
                    "font_name": run.font.name,
                }
            )

        paragraphs.append(
            {
                "text": paragraph.text,
                "level": paragraph.level,
                "alignment": enum_name(paragraph.alignment) if paragraph.alignment is not None else None,
                "runs": runs,
            }
        )

    return paragraphs


def extract_table(shape: Any) -> list[list[str]]:
    if not getattr(shape, "has_table", False):
        return []

    table_data: list[list[str]] = []
    for row in shape.table.rows:
        table_data.append([cell.text for cell in row.cells])
    return table_data


def extract_chart(shape: Any) -> dict[str, Any] | None:
    if not getattr(shape, "has_chart", False):
        return None

    chart = shape.chart
    data: dict[str, Any] = {
        "chart_type": enum_name(chart.chart_type),
        "series": [],
    }

    try:
        for series in chart.series:
            data["series"].append(
                {
                    "name": str(series.name),
                    "values": [float(value) if isinstance(value, (int, float)) else value for value in series.values],
                }
            )
    except Exception:
        data["series"] = []

    return data


def shape_descriptor(shape: Any, z_order: int, group_path: str = "") -> dict[str, Any]:
    descriptor = {
        "shape_id": getattr(shape, "shape_id", None),
        "name": getattr(shape, "name", ""),
        "shape_type": enum_name(getattr(shape, "shape_type", "")),
        "z_order": z_order,
        "group_path": group_path,
        "alt_text": shape_alt_text(shape),
        "geometry": geometry(shape),
    }

    if getattr(shape, "has_text_frame", False):
        descriptor["text"] = shape.text
        descriptor["paragraphs"] = paragraphs_from_text_frame(shape)

    if getattr(shape, "has_table", False):
        descriptor["table"] = extract_table(shape)

    chart = extract_chart(shape)
    if chart:
        descriptor["chart"] = chart

    return descriptor


def iter_shapes(shapes: Any, group_path: str = ""):
    for z_order, shape in enumerate(shapes):
        yield shape, z_order, group_path

        child_shapes = getattr(shape, "shapes", None)
        if child_shapes is not None:
            next_group_path = f"{group_path}/{shape.name}" if group_path else str(shape.name)
            yield from iter_shapes(child_shapes, next_group_path)


def write_image(shape: Any, slide_num: int, image_index: int, assets_dir: Path) -> dict[str, Any]:
    image = shape.image
    image_ext = image.ext
    image_name = f"slide{slide_num}_img{image_index}.{image_ext}"
    image_path = assets_dir / image_name
    image_path.write_bytes(image.blob)

    return {
        "path": f"assets/{image_name}",
        "width": int(shape.width),
        "height": int(shape.height),
        "left": int(shape.left),
        "top": int(shape.top),
        "width_px": emu_to_px(shape.width),
        "height_px": emu_to_px(shape.height),
        "left_px": emu_to_px(shape.left),
        "top_px": emu_to_px(shape.top),
        "alt_text": shape_alt_text(shape),
        "content_type": getattr(image, "content_type", None),
    }


def extract_pptx(file_path: Path, output_dir: Path) -> list[dict[str, Any]]:
    Presentation, MSO_SHAPE_TYPE = load_pptx_modules()

    prs = Presentation(str(file_path))
    slides_data: list[dict[str, Any]] = []

    output_dir.mkdir(parents=True, exist_ok=True)
    assets_dir = output_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)

    slide_width = int(prs.slide_width)
    slide_height = int(prs.slide_height)

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_data: dict[str, Any] = {
            "number": slide_num,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
            "slide_size": {
                "width": slide_width,
                "height": slide_height,
                "width_px": emu_to_px(slide_width),
                "height_px": emu_to_px(slide_height),
                "aspect_ratio": round(slide_width / slide_height, 4) if slide_height else None,
            },
            "shapes": [],
            "tables": [],
            "charts": [],
        }

        image_index = 0

        for shape, z_order, group_path in iter_shapes(slide.shapes):
            descriptor = shape_descriptor(shape, z_order, group_path)
            slide_data["shapes"].append(descriptor)

            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if shape == slide.shapes.title:
                    slide_data["title"] = text
                elif text:
                    slide_data["content"].append(
                        {
                            "type": "text",
                            "content": shape.text,
                            "paragraphs": descriptor.get("paragraphs", []),
                            "geometry": descriptor["geometry"],
                            "z_order": z_order,
                            "name": descriptor["name"],
                        }
                    )

            if getattr(shape, "has_table", False):
                slide_data["tables"].append(
                    {
                        "rows": extract_table(shape),
                        "geometry": descriptor["geometry"],
                        "z_order": z_order,
                        "name": descriptor["name"],
                    }
                )

            chart = extract_chart(shape)
            if chart:
                chart["geometry"] = descriptor["geometry"]
                chart["z_order"] = z_order
                chart["name"] = descriptor["name"]
                slide_data["charts"].append(chart)

            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                image_index += 1
                slide_data["images"].append(write_image(shape, slide_num, image_index, assets_dir))

        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            slide_data["notes"] = notes_frame.text

        slides_data.append(slide_data)

    return slides_data


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract PowerPoint slide content into JSON and image assets.")
    parser.add_argument("input_pptx", help="Path to the .pptx file")
    parser.add_argument("output_dir", nargs="?", default=".", help="Output directory, default: current directory")
    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    input_file = Path(args.input_pptx).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    if not input_file.exists():
        print(f"Input file not found: {input_file}", file=sys.stderr)
        return 1
    if input_file.suffix.lower() != ".pptx":
        print(f"Input file must be a .pptx file: {input_file}", file=sys.stderr)
        return 1

    slides = extract_pptx(input_file, output_dir)

    output_path = output_dir / "extracted-slides.json"
    output_path.write_text(json.dumps(slides, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Extracted {len(slides)} slides to {output_path}")
    for slide in slides:
        img_count = len(slide["images"])
        shape_count = len(slide["shapes"])
        table_count = len(slide["tables"])
        chart_count = len(slide["charts"])
        title = slide["title"] or "(no title)"
        print(
            f"  Slide {slide['number']}: {title} — "
            f"{img_count} image(s), {shape_count} shape(s), "
            f"{table_count} table(s), {chart_count} chart(s)"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
